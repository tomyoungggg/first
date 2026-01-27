import pandas as pd
from datetime import datetime
import requests
from config import METABASE_URL, METABASE_API_KEY
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
import numpy as np

print("=" * 80)
print("SALES RETENTION ANALYSIS - NET LOSSES & AFFILIATE PAYOUTS")
print("=" * 80)

ACTUALS_SHEET_ID = '1-kbUBDVdLIXf0bbG4C_c7JWbyipGQtCdWjW7R7agiaM'
BUDGET_SUMMARY_SHEET_ID = '1PY4MPmPXptZ3p4kJOmtS-mP35AfMq2DAvFG56CA9ffQ'
CREDENTIALS_FILE = r'C:\Users\Tom Young\OneDrive\Desktop\credentials.json'
CARD_ID = 13018  # Net Losses AND Affiliate Payouts

metabase_headers = {
    "X-API-KEY": METABASE_API_KEY,
    "Content-Type": "application/json"
}

# Step 1: Fetch ACTUALS Sales spend (completed months)
print("\n1. Fetching ACTUALS Sales spend (completed months)...")

try:
    credentials = Credentials.from_service_account_file(
        CREDENTIALS_FILE,
        scopes=['https://www.googleapis.com/auth/spreadsheets.readonly']
    )
    sheets_service = build('sheets', 'v4', credentials=credentials)

    result = sheets_service.spreadsheets().values().get(
        spreadsheetId=ACTUALS_SHEET_ID,
        range='Sheet1!B2:F100'
    ).execute()

    values = result.get('values', [])

    sales_spend = {}

    for row in values:
        if len(row) < 5:
            continue

        month_str = row[0]
        sales_str = row[2]

        if not month_str or not sales_str:
            continue

        try:
            month_date = pd.to_datetime(month_str, errors='coerce')
            if pd.isna(month_date):
                continue

            month_key = month_date.strftime('%Y-%m')
            sales_val = float(str(sales_str).replace('$', '').replace(',', '').strip())
            sales_spend[month_key] = sales_val

        except:
            continue

    print(f"   ✓ Retrieved actuals for {len(sales_spend)} months")

except Exception as e:
    print(f"   ❌ Error: {e}")
    sales_spend = {}

# Step 2: Fetch BUDGET for current incomplete month (Jan 2026)
print("\n2. Fetching BUDGET spend for current month from Budget Summary...")

try:
    result = sheets_service.spreadsheets().values().batchGet(
        spreadsheetId=BUDGET_SUMMARY_SHEET_ID,
        ranges=[
            'Budget Summary!G27',  # Creator referral
            'Budget Summary!G34',  # Sales Spend
            'Budget Summary!G42',  # AM Spend
        ]
    ).execute()

    value_ranges = result.get('valueRanges', [])

    if len(value_ranges) >= 3:
        creator_referral_val = 0
        sales_spend_val = 0
        am_spend_val = 0

        if value_ranges[0].get('values'):
            creator_referral_val = float(str(value_ranges[0]['values'][0][0]).replace('$', '').replace(',', '').strip())

        if value_ranges[1].get('values'):
            sales_spend_val = float(str(value_ranges[1]['values'][0][0]).replace('$', '').replace(',', '').strip())

        if value_ranges[2].get('values'):
            am_spend_val = float(str(value_ranges[2]['values'][0][0]).replace('$', '').replace(',', '').strip())

        # Calculate Sales spend for Jan 2026
        jan_2026_sales = sales_spend_val + am_spend_val - creator_referral_val
        sales_spend['2026-01'] = jan_2026_sales

        print(f"   ✓ Jan 2026 Budget:")
        print(f"     Sales: ${jan_2026_sales:,.0f} (Sales ${sales_spend_val:,.0f} + AM ${am_spend_val:,.0f} - Creator Ref ${creator_referral_val:,.0f})")
    else:
        print(f"   ⚠ Could not read budget values")

except Exception as e:
    print(f"   ❌ Error: {e}")
    import traceback
    traceback.print_exc()

print(f"\n   ✓ Total months with spend data: {len(sales_spend)}")

# Step 3: Fetch Metabase data
print("\n3. Fetching data from Metabase card 13018 (Net Losses AND Affiliate Payouts)...")

try:
    response = requests.post(
        f"{METABASE_URL}/api/card/{CARD_ID}/query/json",
        headers=metabase_headers
    )
    response.raise_for_status()
    data = response.json()

    if isinstance(data, list) and len(data) > 0:
        df = pd.DataFrame(data)
        print(f"   ✓ Retrieved {len(df)} rows, {len(df.columns)} columns")
    else:
        print(f"   ❌ No data returned")
        exit(1)

except Exception as e:
    print(f"   ❌ Error: {e}")
    exit(1)

# Step 4: Identify columns
cohort_col = None
period_col = None
value_col = None

for col in df.columns:
    col_lower = str(col).lower()
    if 'first transaction month' in col_lower or 'first_transaction_month' in col_lower:
        cohort_col = col
    elif 'period' in col_lower:
        period_col = col
    elif 'contribution' in col_lower or 'profit' in col_lower:
        value_col = col

if not all([cohort_col, period_col, value_col]):
    print(f"\n❌ Missing columns")
    exit(1)

print(f"   ✓ Using columns: {cohort_col}, {period_col}, {value_col}")

# Step 5: Create pivot
print("\n4. Creating analysis...")

pivot = df.pivot_table(
    values=value_col,
    index=cohort_col,
    columns=period_col,
    aggfunc='sum',
    fill_value=0
)

print(f"   ✓ Pivot: {pivot.shape[0]} cohorts x {pivot.shape[1]} periods")

# Create cumulative
cumulative = pivot.copy()
current_date = datetime.now()

for cohort in cumulative.index:
    cohort_str = str(cohort)

    try:
        if 'T' in cohort_str:
            cohort_date = pd.to_datetime(cohort_str.split('T')[0])
        else:
            cohort_date = pd.to_datetime(cohort_str)

        months_elapsed = (current_date.year - cohort_date.year) * 12 + (current_date.month - cohort_date.month) + 1

        row_values = pivot.loc[cohort].values
        cumulative_values = []
        running_sum = 0

        for period_idx, val in enumerate(row_values):
            period_num = period_idx + 1
            if period_num <= months_elapsed:
                running_sum += val
                cumulative_values.append(running_sum)
            else:
                cumulative_values.append(np.nan)

        cumulative.loc[cohort] = cumulative_values

    except:
        cumulative.loc[cohort] = pivot.loc[cohort].cumsum()

print(f"   ✓ Cumulative created")

# Match spend to cohorts
spend_matched = {}

print(f"\n   Matching Sales spend to cohorts:")
for cohort in pivot.index:
    cohort_str = str(cohort)

    try:
        if 'T' in cohort_str:
            cohort_date = pd.to_datetime(cohort_str.split('T')[0])
        else:
            cohort_date = pd.to_datetime(cohort_str)

        month_key = cohort_date.strftime('%Y-%m')

        if month_key in sales_spend:
            spend_matched[cohort] = sales_spend[month_key]
            source = "BUDGET" if month_key == "2026-01" else "ACTUALS"
            print(f"   ✓ {cohort_str[:10]} → {month_key} → ${spend_matched[cohort]:,.0f} ({source})")
        else:
            print(f"   ✗ {cohort_str[:10]} → {month_key} (not found)")

    except Exception as e:
        print(f"   ✗ {cohort_str[:10]} - error: {e}")

spend_series = pd.Series(spend_matched, name='Sales Spend')
print(f"\n   ✓ Matched spend for {len(spend_matched)}/{len(pivot.index)} cohorts")

# Calculate payback ratios
payback = cumulative.copy()
for cohort in payback.index:
    if cohort in spend_matched and spend_matched[cohort] > 0:
        spend = spend_matched[cohort]
        payback.loc[cohort] = cumulative.loc[cohort] / spend
    else:
        payback.loc[cohort] = np.nan

print(f"   ✓ Payback ratios calculated")

# Step 6: Write to Excel - ONE sheet with cumulative on top, payback below
output_file = f'Sales_Retention_Analysis_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

print(f"\n5. Creating Excel file: {output_file}")

with pd.ExcelWriter(output_file, engine='openpyxl') as writer:

    # Write cumulative at top
    cumulative.to_excel(writer, sheet_name='Analysis', startrow=0)

    # Write payback below with spacing
    payback_start_row = cumulative.shape[0] + 3

    # Add spend column to payback
    payback_with_spend = payback.copy()
    payback_with_spend.insert(0, 'Sales Spend', spend_series)
    payback_with_spend.to_excel(writer, sheet_name='Analysis', startrow=payback_start_row)

    # Format
    from openpyxl.styles import numbers, Font

    ws = writer.sheets['Analysis']

    # Format cumulative section (currency)
    for row in range(2, cumulative.shape[0] + 2):
        for col in range(2, cumulative.shape[1] + 2):
            cell = ws.cell(row=row, column=col)
            cell.number_format = '$#,##0'

    # Add "Payback Analysis" header
    header_cell = ws.cell(row=payback_start_row, column=1)
    header_cell.value = "Payback Analysis"
    header_cell.font = Font(bold=True)

    # Format payback section - spend column as currency (column 2)
    for row in range(payback_start_row + 2, payback_start_row + payback_with_spend.shape[0] + 2):
        cell = ws.cell(row=row, column=2)
        cell.number_format = '$#,##0'

    # Format payback ratios as percentage (columns 3+)
    for row in range(payback_start_row + 2, payback_start_row + payback_with_spend.shape[0] + 2):
        for col in range(3, payback_with_spend.shape[1] + 2):
            cell = ws.cell(row=row, column=col)
            cell.number_format = '0.00%'

    print(f"   ✓ Exported to single 'Analysis' sheet")

print(f"\n✅ Complete!")
print(f"\nFile: {output_file}")
print(f"\nSheet contains:")
print(f"  - Cumulative profit tracking (top)")
print(f"  - Payback % analysis with Sales spend (below)")
print("\n" + "=" * 80)
